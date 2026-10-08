#!/usr/bin/env python3
"""Tab XNBH — bộ nước rửa kính (BNR) & xác nhận bảo hành / NPS.

Nguồn: file BNR*.xlsx đặt cùng cấp với thư mục repo (C:\\CRM\\CRM\\BNR T8.26.xlsx), sheet "BNR".
Ghi kết quả vào khoá "xnbh" của .work/dash.json. Chỉ ghi số tổng hợp — KHÔNG đưa Mã KH (SĐT) vào.

Chỉ số chính: tỷ lệ XNBH trên KH nhận BNR = xnb / kh
  xnb = dòng XNBH=Yes có SL KH   (khách nhận BNR và đã xác nhận bảo hành)
  xno = dòng XNBH=Yes không SL KH (xác nhận bảo hành nhưng không nhận BNR)

Script tự kiểm tra hai thứ và ghi vào dash.json để giao diện cảnh báo:
  mode[m]  "flag"  = cột XNBH phân biệt được đã/chưa xác nhận (bình thường)
           "grant" = cột XNBH chỉ lặp lại "có nhận BNR hay không" -> không dùng để đo chuyển đổi được
  mism[m]  số dòng mà cột XNBH và cột NPS không khớp nhau (bình thường = 0)
"""
import os, sys, glob, json, collections, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
WORK = os.path.join(ROOT, ".work")
DASH = os.path.join(WORK, "dash.json")

def find_src():
    pats = [os.path.join(os.path.dirname(ROOT), "BNR*.xlsx"), os.path.join(ROOT, "BNR*.xlsx")]
    for p in pats:
        hits = [f for f in sorted(glob.glob(p)) if not os.path.basename(f).startswith("~$")]
        if hits: return hits[-1]
    return None

def main():
    src = find_src()
    if not src:
        print("XNBH: không tìm thấy file BNR*.xlsx — bỏ qua"); return
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if "BNR" not in wb.sheetnames:
        print("XNBH: file", os.path.basename(src), "không có sheet BNR — bỏ qua"); return
    ws = wb["BNR"]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h or "").strip() for h in next(it)]
    I = {h: i for i, h in enumerate(hdr)}
    need = ["CH", "SL tặng", "SL KH", "NPS", "XNBH", "Month"]
    miss = [c for c in need if c not in I]
    if miss:
        print("XNBH: thiếu cột", miss, "— bỏ qua"); return
    g = lambda r, c: r[I[c]] if I.get(c) is not None and I[c] < len(r) else None

    agg = collections.defaultdict(lambda: {"rows": 0, "kh": 0, "sets": 0, "xnb": 0, "xno": 0, "npsb": 0, "npso": 0, "mism": 0, "xor": 0})
    prof = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    nrow = 0
    for r in it:
        if not any(v is not None for v in r): continue
        kho = str(g(r, "CH") or "").strip()
        mo = g(r, "Month")
        if not kho or mo is None: continue
        m = mo.strftime("%Y-%m") if isinstance(mo, (datetime.datetime, datetime.date)) else str(mo)[:7]
        nrow += 1
        kh = int(g(r, "SL KH") or 0)
        st = int(g(r, "SL tặng") or 0)
        yes = str(g(r, "XNBH") or "").strip().lower() == "yes"
        nps = (g(r, "NPS") or 0) and True
        nps = bool(nps)
        a = agg[(m, kho)]
        a["rows"] += 1; a["kh"] += kh; a["sets"] += st
        if yes:
            if kh > 0: a["xnb"] += 1
            else: a["xno"] += 1
        if nps:
            if kh > 0: a["npsb"] += 1
            else: a["npso"] += 1
        if yes != nps: a["mism"] += 1          # cột XNBH và NPS lệch nhau
        if yes != (kh > 0): a["xor"] += 1      # XNBH khác với "có nhận BNR"
        for c in ("Tier", "Khu vực", "AM"):
            v = g(r, c)
            if v: prof[kho][c][str(v).strip()] += 1

    # model của cửa hàng lấy từ dash.json để dùng chung bộ lọc Model
    d = json.load(open(DASH, encoding="utf-8"))
    model = {s["kho"]: s.get("model", "") for s in d.get("stores", [])}

    rows = [dict(month=m, kho=kho, model=model.get(kho, ""), **v) for (m, kho), v in sorted(agg.items())]
    stores = {k: {"tier": c["Tier"].most_common(1)[0][0] if c["Tier"] else "",
                  "khuvuc": c["Khu vực"].most_common(1)[0][0] if c["Khu vực"] else "",
                  "am": c["AM"].most_common(1)[0][0] if c["AM"] else ""} for k, c in prof.items()}
    months = sorted({r["month"] for r in rows})

    # Ngữ nghĩa cột XNBH theo từng tháng.
    # Dò bằng số dòng mà (XNBH=Yes) khác với (có nhận BNR): gần 0 nghĩa là cột XNBH
    # chỉ lặp lại việc có nhận BNR hay không -> không đo được chuyển đổi.
    mode, mism = {}, {}
    for m in months:
        rs = [r for r in rows if r["month"] == m]
        nr = sum(r["rows"] for r in rs); xor = sum(r["xor"] for r in rs)
        mode[m] = "grant" if nr and xor <= max(2, 0.005 * nr) else "flag"
        mism[m] = sum(r["mism"] for r in rs)

    d["xnbh"] = {"rows": rows, "months": months, "stores": stores, "mode": mode, "mism": mism,
                 "src": os.path.basename(src), "nrow": nrow,
                 "generated": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")}
    json.dump(d, open(DASH, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    tot = collections.Counter()
    for r in rows:
        for k in ("kh", "sets", "xnb", "xno", "npsb", "npso", "mism"): tot[k] += r[k]
    print(f"XNBH ok: {os.path.basename(src)} · {nrow} dòng · {len(months)} tháng ({months[0]}→{months[-1]}) · "
          f"{len(set(r['kho'] for r in rows))} CH · KH nhận BNR {tot['kh']} · XNBH kèm BNR {tot['xnb']} · XNBH lẻ {tot['xno']}")
    print("   cột XNBH theo tháng:", mode, "| số dòng XNBH lệch NPS:", mism)

if __name__ == "__main__":
    main()
